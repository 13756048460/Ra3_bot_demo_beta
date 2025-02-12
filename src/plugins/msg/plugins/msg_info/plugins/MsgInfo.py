from nonebot import on_keyword, on_fullmatch
from nonebot.adapters.onebot.v11 import MessageSegment, event, Bot, GroupMessageEvent, bot, Message
from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook
from nonebot import require
require("nonebot_plugin_localstore")
import nonebot_plugin_localstore as store
import os
import asyncio
import nonebot

config = nonebot.get_driver().config
files = config.files

class MsgInfo:
    def __init__(self, info, msg,mag_type):
        self.info = info
        self.msg = msg
        self.mag_type = mag_type
    def __str__(self):
        return f"{self.info} {self.msg} {self.mag_type}"

async def read_xlsx(file_name):
    """
    :param file_name:
    :return:
    """
    wb = await asyncio.to_thread(load_workbook, file_name)
    sheet = wb.active
    wb.close()
    for row in sheet.iter_rows(min_row=wb.active.min_row + 1, max_row=wb.active.max_row):
        msg_info = MsgInfo(row[0].value, row[1].value,row[2].value)
        yield msg_info

def check_file_exists(folder_path, file_name):
    """

    :param folder_path:
    :param file_name:
    :return:
    """
    file_path = os.path.join(folder_path, file_name)
    return os.path.exists(file_path)


def split_msg(info):
    """

    :param info:
    :return:
    """
    result = [part for part in info.split("#") if part]
    set_from_arr = set(result) #?????????????????????????????????
    return set_from_arr
def if_msg(msg_info, user_info):
    """
    :param msg_info:
    :param user_info:
    :return:
    """
    split_msg_info = split_msg(msg_info.info)
    if user_info in split_msg_info:
        return msg_info.msg
    return None


def read_info(file_name):
    """
    :param file_name:
    :return:
    """
    arr = []
    set_from_arr = set()
    wb = load_workbook(file_name)
    sheet = wb.active
    wb.close()
    for row in sheet.iter_rows(min_row=wb.active.min_row + 1):
        arr.append(row[0].value)
    for msg_info in arr:
        result = [part for part in msg_info.split("#") if part]
        set_from_arr.update(set(result))
    return set_from_arr

ms = set(read_info(f'{files}/xlsx/msg.xlsx'))
info = on_keyword(keywords=ms)

@info.handle()
async def _(bot: Bot, event: GroupMessageEvent):
    user_info = str(event.get_plaintext())
    async for msg_info in read_xlsx(f'{files}/xlsx/msg.xlsx'):
        res = if_msg(msg_info, user_info)
        if res is not None:
            msg = res
            #msg = str(msg_info.mag_type)
            break
    if msg_info.mag_type is not None:
        messages = [
            {
                "type": "node",
                "data": {
                    "name": bot.self_id,
                    "uin": bot.self_id,
                    "content": [MessageSegment.text(msg)],
                },
            },
        ]
        res_id = await bot.call_api("send_forward_msg", messages=messages)
        await bot.send_group_msg(group_id=event.group_id, message=MessageSegment.forward(res_id))
    else:
        await info.finish(MessageSegment.text(msg))